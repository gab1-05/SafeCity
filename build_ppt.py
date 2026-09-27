"""Build SafeCity SDP presentation (16:9 .pptx) — StayMate format + DevOps deployment thread."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY = RGBColor(0x0F, 0x2A, 0x3C)
TEAL = RGBColor(0x1B, 0x6F, 0x8B)
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
LIGHT = RGBColor(0xF4, 0xF7, 0xF9)
DARK_TXT = RGBColor(0x1E, 0x2B, 0x3A)
GREY = RGBColor(0x5B, 0x6B, 0x7C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 17


def header(slide, num):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.62))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tx = bar.text_frame; tx.word_wrap = True
    p = tx.paragraphs[0]
    p.text = "SDP PROJECT  •  SafeCity — Smart City Incident Management & Civic Response Platform"
    p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = WHITE
    tag = slide.shapes.add_textbox(prs.slide_width - Inches(1.5), Inches(0.1), Inches(1.3), Inches(0.42))
    tf = tag.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    p.text = f"{num} / {TOTAL}"; p.font.size = Pt(12); p.font.color.rgb = WHITE
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.62), prs.slide_width, Pt(4))
    line.fill.solid(); line.fill.fore_color.rgb = TEAL
    line.line.fill.background()


def title_block(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12.1), Inches(1.1))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = title
    p.font.size = Pt(34); p.font.bold = True; p.font.color.rgb = NAVY
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.6), Inches(1.85), Inches(12.1), Inches(0.6))
        tf2 = box2.text_frame; tf2.word_wrap = True
        p = tf2.paragraphs[0]; p.text = subtitle
        p.font.size = Pt(16); p.font.color.rgb = GREY


def bullets(slide, items, top=2.6, size=17, gap=0.62, bold_prefix=True):
    for i, it in enumerate(items):
        box = slide.shapes.add_textbox(Inches(0.7), Inches(top + i * gap), Inches(11.9), Inches(0.6))
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        if bold_prefix and "—" in it:
            head, rest = it.split("—", 1)
            r = p.add_run(); r.text = "▸ " + head.strip() + " — "
            r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = TEAL
            r2 = p.add_run(); r2.text = rest.strip()
            r2.font.size = Pt(size); r2.font.color.rgb = DARK_TXT
        else:
            r = p.add_run(); r.text = "▸ " + it
            r.font.size = Pt(size); r.font.color.rgb = DARK_TXT


def grid_boxes(slide, items, top=2.7, cols=3, box_h=1.75, box_w=3.85, hgap=0.25, left0=0.64):
    for i, (h, b) in enumerate(items):
        r, c = divmod(i, cols)
        x = Inches(left0 + c * (box_w + hgap)); y = Inches(top + r * (box_h + 0.25))
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(box_w), Inches(box_h))
        shp.fill.solid(); shp.fill.fore_color.rgb = WHITE
        shp.line.color.rgb = TEAL; shp.line.width = Pt(1.5)
        tf = shp.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = h
        run.font.size = Pt(16); run.font.bold = True; run.font.color.rgb = TEAL
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run(); run2.text = b
        run2.font.size = Pt(13); run2.font.color.rgb = DARK_TXT


def table_slide(slide, rows, top=2.6, col_widths=None, font=12):
    n_r, n_c = len(rows), len(rows[0])
    if col_widths is None:
        col_widths = [Inches(12.1 / n_c)] * n_c
    tbl_shape = slide.shapes.add_table(n_r, n_c, Inches(0.6), Inches(top),
                                       sum(col_widths), Inches(0.5 + n_r * 0.42))
    tbl = tbl_shape.table
    for j, w in enumerate(col_widths):
        tbl.columns[j].width = w
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run(); run.text = val
            run.font.size = Pt(font + (1 if i == 0 else 0))
            run.font.bold = (i == 0)
            run.font.color.rgb = WHITE if i == 0 else DARK_TXT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (WHITE if i % 2 else LIGHT)


def note_box(slide, text, top=5.35):
    ph = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(top), Inches(12.1), Inches(1.1))
    ph.fill.solid(); ph.fill.fore_color.rgb = LIGHT
    ph.line.color.rgb = ORANGE; ph.line.width = Pt(2)
    tf = ph.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = ORANGE


# 1 — Title
s = prs.slides.add_slide(BLANK)
header(s, 1)
t = s.shapes.add_textbox(Inches(0.6), Inches(1.2), Inches(12.1), Inches(1.2))
tf = t.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "SafeCity"
r.font.size = Pt(54); r.font.bold = True; r.font.color.rgb = NAVY
t2 = s.shapes.add_textbox(Inches(1.3), Inches(2.35), Inches(10.7), Inches(0.7))
tf = t2.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Smart City Incident Management & Civic Response Platform"
r.font.size = Pt(22); r.font.color.rgb = TEAL
t3 = s.shapes.add_textbox(Inches(1.3), Inches(3.1), Inches(10.7), Inches(0.5))
tf = t3.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "\u201CReport it. Track it. Resolve it.\u201D — Accountable, SLA-driven & Citizen-visible"
r.font.size = Pt(15); r.font.italic = True; r.font.color.rgb = GREY
for i, txt in enumerate(["[ Group No: ____ ]", "[ Member 1 — Name (Roll) ]   [ Member 2 — Name (Roll) ]   [ Member 3 — Name (Roll) ]   [ Member 4 — Name (Roll) ]", "[ Under Guidance of: Prof. ____ ]   •   The Bombay Salesian Society"]):
    b = s.shapes.add_textbox(Inches(1.3), Inches(4.15 + i * 0.55), Inches(10.7), Inches(0.5))
    tf = b.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = txt
    r.font.size = Pt(14); r.font.color.rgb = DARK_TXT
    if i == 0:
        r.font.bold = True

# 2 — Problem
s = prs.slides.add_slide(BLANK)
header(s, 2)
title_block(s, "Problem Statement", "Civic complaints run on fragmented channels — SafeCity makes them one accountable workflow")
bullets(s, [
    "No single intake — helpline calls, WhatsApp forwards, social posts and walk-ins stay invisible to each other; one pothole gets reported five times and fixed twice",
    "No shared status vocabulary — \u201Cforwarded\u201D is indistinguishable from \u201Cfixed\u201D; citizens cannot tell whether a report is being worked on",
    "No accountability trail — once a complaint leaves the citizen's hands there is no record of who acted, when, or why it closed",
    "No SLA discipline — without deadlines, low-visibility issues (streetlights, garbage) starve while the loudest escalation wins",
    "No prioritisation signal — a fallen live wire and a faded road marking enter the same queue in arrival order",
    "No feedback loop — reporting appears to accomplish nothing, so citizens stop reporting and the city loses its own signal",
])

# 3 — Objectives
s = prs.slides.add_slide(BLANK)
header(s, 3)
title_block(s, "Objectives", "O1–O10 from OVERVIEW.md, condensed to six outcomes")
grid_boxes(s, [
    ("Single intake", "One reporting wizard: category, location, evidence, severity + anonymity"),
    ("Correct routing", "Every category declares an owning department; auto + workload-aware assign"),
    ("Enforced lifecycle", "Declarative transition matrix — illegal moves rejected, never silently ignored"),
    ("SLA discipline", "Per-(category, severity) deadlines; Celery sweep flags every breach"),
    ("Auditability", "Append-only audit row in the same DB transaction as each state change"),
    ("Visibility + privacy", "Public tracking + reference lookup with PII stripped and coords jittered"),
])

# 4 — Scope
s = prs.slides.add_slide(BLANK)
header(s, 4)
title_block(s, "Scope", "What SafeCity covers — and honestly excludes")
table_slide(s, [
    ["Area", "In scope", "Out of scope"],
    ["Users", "Citizen register/auth, profiles, privacy, consents, deletion requests", "Native mobile apps (SPA is responsive to 360px)"],
    ["Workflow", "Verify → assign → work → resolve → confirm/close + reject, duplicate/merge, escalate, reopen", "Official emergency dispatch (escalated to a human, never auto-dispatched)"],
    ["Ops", "SLA computation + breach sweeps, notifications (in-app/WS/email), role dashboards, analytics + CSV", "Government API integration (all externals are mockable adapters)"],
    ["Platform", "Docker Compose, Helm, Terraform, CI (lint/typecheck/test), audit tooling", "Auto-applying AWS infra, payments/billing"],
], col_widths=[Inches(1.4), Inches(5.35), Inches(5.35)])

# 5 — Survey placeholder
s = prs.slides.add_slide(BLANK)
header(s, 5)
title_block(s, "Citizen Survey (add your Google Form results here)", "Suggested questions below — replace with your charts before Review 1")
bullets(s, [
    "Q1 — How do you currently report civic issues? (helpline / WhatsApp / social / walk-in / app)",
    "Q2 — Did you ever learn what happened to your complaint? (yes / no / partially)",
    "Q3 — Biggest pain point? (no tracking / no action / slow response / rude process)",
    "Q4 — Would you use a single app with live tracking + SLA deadlines? (1–5 scale)",
], size=16)
note_box(s, "PLACEHOLDER — paste 2–3 survey charts (response mix, tracking gap, willingness) in this box")

# 6 — Existing systems
s = prs.slides.add_slide(BLANK)
header(s, 6)
title_block(s, "Existing System Survey", "What cities use today — and the gap SafeCity closes")
table_slide(s, [
    ["#", "System", "What it does", "Gap SafeCity fills"],
    ["1", "NYC/SF 311 + helplines", "Phone/app intake for non-emergency city services", "No enforced workflow, SLA or audit visible to citizens"],
    ["2", "BMC/MCGM portals + WhatsApp", "Departmental inboxes and forwards", "Fragmented channels; same issue reported many times, worked twice"],
    ["3", "IChangeMyCity (Janaagraha)", "Civic reporting with community votes", "Weak routing/SLA/escalation and no immutable accountability trail"],
    ["4", "Fix-It (AI grievance pilots)", "AI-assisted municipal complaint filing with photos", "AI acts without human-gated decisions or privacy-safe public views"],
], col_widths=[Inches(0.5), Inches(2.4), Inches(4.4), Inches(4.8)])

# 7 — Literature (extended with DevOps refs)
s = prs.slides.add_slide(BLANK)
header(s, 7)
title_block(s, "Literature Survey", "Evidence behind the design — product + DevOps")
table_slide(s, [
    ["#", "Work", "Finding for SafeCity"],
    ["5", "311 & smart-city governance bias (Sustainable Cities & Society, 2020)", "Validates 311 intake; warns of reporting gaps — hence anonymity + ward analytics"],
    ["6", "NIST RBAC model (Ferraiolo & Kuhn) + RBAC on the Web", "Server-enforced role matrix; no permission decided client-side"],
    ["7", "Jenkins automation efficiency for containers (Sensors, 2024)", "Quantifies CI-driven release gains — our build→test→deploy template"],
    ["8", "Full-stack CI/CD deployment toolchains (IJACTE, 2026)", "Git→Jenkins→SonarQube→Trivy→Docker→K8s integration justification"],
    ["9", "DevSecOps CI/CD on AWS review (IJERT, 2026)", "Multi-layer security scanning + secure microservice deployment pattern"],
    ["10", "Explainable multilingual civic-complaint AI (IEEE WConf.)", "Advisory-only AI: scrubbed input, human-gated decisions"],
], col_widths=[Inches(0.5), Inches(4.3), Inches(7.3)], font=11)

# 8 — System architecture
s = prs.slides.add_slide(BLANK)
header(s, 8)
title_block(s, "System Architecture", "React SPA → Django API → Postgres/Redis/Celery — one container per concern")
grid_boxes(s, [
    ("FRONTEND", "React 18 + TS + Vite\nTailwind, TanStack Query\nLeaflet maps, WS client"),
    ("BACKEND", "Django 5 + DRF\nJWT + RBAC + transition\nmatrix + SLA engine"),
    ("DATA & JOBS", "PostgreSQL 16\nRedis + Celery/beat\nMinIO (optional S3)"),
], top=2.7, cols=3, box_h=2.6)
flow = s.shapes.add_textbox(Inches(0.6), Inches(5.5), Inches(12.1), Inches(0.6))
tf = flow.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Citizen wizard → REST /api/v1 + WS /ws/notifications → Postgres (audit in same txn) → Celery sweeps → dashboards/analytics"
r.font.size = Pt(14); r.font.color.rgb = TEAL; r.font.bold = True

# 9 — DevOps pipeline
s = prs.slides.add_slide(BLANK)
header(s, 9)
title_block(s, "DevOps Pipeline", "Every merge verified: lint → test → scan → build → validate → ship")
steps = ["Git + GitHub\n(branch + PR)", "CI gates\n(ruff, tsc, eslint\npytest, vitest)", "Security\n(Trivy images\nSonarCloud)", "Docker\n(build + push\nGHCR)", "Validate\n(helm lint/template\nkubeconform, tf fmt)", "Ship + observe\n(compose / kind\nhealth + readiness)"]
for i, st in enumerate(steps):
    x = 0.6 + i * 2.1
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.8), Inches(1.9), Inches(1.9))
    shp.fill.solid(); shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = TEAL; shp.line.width = Pt(1.5)
    tf = shp.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = st
    r.font.size = Pt(13); r.font.color.rgb = DARK_TXT
    if i < len(steps) - 1:
        ar = s.shapes.add_textbox(Inches(x + 1.92), Inches(3.55), Inches(0.25), Inches(0.4))
        p = ar.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = "›"
        r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = ORANGE
note = s.shapes.add_textbox(Inches(0.6), Inches(5.1), Inches(12.1), Inches(1.2))
tf = note.text_frame; tf.word_wrap = True
for line in ["Zero-credential start: make docker-up && make migrate && make seed — no cloud accounts needed.",
             "Jenkinsfile (optional, rubric-only) mirrors these stages on local Docker agents."]:
    p = tf.add_paragraph(); p.space_after = Pt(4)
    r = p.add_run(); r.text = "▸ " + line
    r.font.size = Pt(14); r.font.color.rgb = DARK_TXT

# 10 — Tool justification
s = prs.slides.add_slide(BLANK)
header(s, 10)
title_block(s, "Tool Justification (free miniproject)", "Why each choice — and the costlier alternative rejected")
table_slide(s, [
    ["Layer", "Choice", "Why (free) / rejected alternative"],
    ["CI", "GitHub Actions (free, OIDC to AWS)", "No servers to host; Jenkins self-host only if rubric names it"],
    ["Registry", "GHCR (free)", "ECR bills storage + needs AWS account wiring"],
    ["K8s", "kind / minikube + in-cluster PG/Redis", "EKS control plane alone ~$73/mo; chart needs dbHost/redisUrl"],
    ["IaC", "Terraform plan-only (fmt/validate/plan)", "Proves design at $0; apply starts ~$185/mo dev billing"],
    ["Observe", "health/readiness + optional Prom/Grafana on kind", "CloudWatch exists only after paid apply"],
], col_widths=[Inches(1.3), Inches(3.4), Inches(7.4)])

# 11 — Deployment architecture
s = prs.slides.add_slide(BLANK)
header(s, 11)
title_block(s, "Deployment Architecture", "What runs free today vs the AWS design (plan-only)")
grid_boxes(s, [
    ("PATH A — LIVE ($0)", "Docker Compose: db, redis, minio\nbackend :8000→18081, frontend :5174\nmigrate + seed, health-gated startup"),
    ("PATH B — K8s PROOF ($0)", "kind + Helm chart (dev overlay)\nBitnami PG/Redis in-cluster\nIngress /api, /ws, /admin routing"),
    ("PATH C — AWS (PLAN ONLY)", "VPC → EKS → RDS + ElastiCache\nS3 media, ALB + ACM, IRSA\nSecrets Manager, CloudWatch"),
], top=2.7, cols=3, box_h=2.6)
flow = s.shapes.add_textbox(Inches(0.6), Inches(5.6), Inches(12.1), Inches(0.6))
tf = flow.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "A proves it runs  •  B proves it is cloud-native  •  C (terraform plan) proves the AWS design — no apply, no bill"
r.font.size = Pt(14); r.font.color.rgb = TEAL; r.font.bold = True

# 12 — Environments & cost
s = prs.slides.add_slide(BLANK)
header(s, 12)
title_block(s, "Environments & $0 Strategy", "Graded DevOps depth without the cloud bill")
table_slide(s, [
    ["Environment", "How", "Cost"],
    ["Dev (live demo)", "Compose + seed; helm template reviewed, never blind-applied", "$0"],
    ["K8s proof", "kind + Helm dev overlay + in-cluster data services", "$0"],
    ["CI", "Actions: backend-ci, frontend-ci, docker-ci, k8s-validate", "$0 (public repo)"],
    ["AWS design", "Terraform fmt/validate/plan reviewed as code; budget alarm first", "$0 (apply refused by Makefile guard)"],
    ["If applied", "EKS $73 + NAT $41×N + RDS + Redis + ALB ≈ $185/mo dev", "Destroy same day; verify NAT/EIP gone"],
], col_widths=[Inches(1.8), Inches(6.6), Inches(3.7)])

# 13 — Plan (DevOps rows added)
s = prs.slides.add_slide(BLANK)
header(s, 13)
title_block(s, "Project Plan", "Six phases — DevOps is gradable work, not an afterthought")
table_slide(s, [
    ["Phase", "Key deliverables", "Owners"],
    ["1 — Planning", "Problem, objectives, SDG map, literature (6+ refs), architecture, toolchain", "All (survey + arch split)"],
    ["2 — Core build", "DB schema, auth/incident/search APIs, citizen + authority pages", "M1 backend, M2 frontend, M4 DB"],
    ["3 — DevOps setup", "Branching, Dockerfiles, Actions CI (4 workflows), Trivy, kind overlay, TF draft (+Jenkinsfile*)", "M3 lead, M4 tests"],
    ["4 — Deploy", "Compose live; kind + Helm verified; TF plan reviewed; dashboards live", "M3 lead, M1+M2 support"],
    ["5 — Harden", "E2E tests, bugfix, helm/kubeconform clean, perf tune, ports/secrets", "M4 lead, all fix"],
    ["6 — Review prep", "Report, demo rehearsal, deployment walkthrough, contributions", "All"],
], col_widths=[Inches(1.7), Inches(6.9), Inches(3.5)], font=11)

# 14 — Roles
s = prs.slides.add_slide(BLANK)
header(s, 14)
title_block(s, "Member Roles (replace placeholders with your names)", "Four hats — each owns a layer end to end")
table_slide(s, [
    ["Role", "Owns", "Presents in Review 1"],
    ["M1 — Backend lead [Name]", "Django/DRF APIs: auth, incidents, workflow, SLA; integration", "Problem, objectives & architecture"],
    ["M2 — Frontend lead [Name]", "React UI: search/filter, report wizard, dashboards, maps", "Proposed system & UI flow"],
    ["M3 — DevOps [Name]", "CI/CD, Docker, Compose/kind/Helm, TF plan, monitoring", "Pipeline, toolchain & deployment"],
    ["M4 — Data + QA [Name]", "Postgres schema, seed, admin moderation, pytest/vitest in CI", "Related work & team roles"],
], col_widths=[Inches(2.6), Inches(5.5), Inches(4.0)])

# 15 — Refs A
s = prs.slides.add_slide(BLANK)
header(s, 15)
title_block(s, "References — Existing Systems", "IEEE-style; accessed Sep 2026")
bullets(s, [
    "NYC 311 & SF 311 — non-emergency municipal intake platforms, nyc.gov/311; sf311.org",
    "IChangeMyCity (Janaagraha) — civic reporting with community support, ichangemycity.com",
    "BMC/MCGM complaint portals + WhatsApp helplines — departmental civic inboxes",
    "Fix-It: AI-based public complaint management (municipal grievance pilots), Semanticscholar",
], size=14, gap=0.72, bold_prefix=False)

# 16 — Refs B
s = prs.slides.add_slide(BLANK)
header(s, 16)
title_block(s, "References — Literature", "Peer-reviewed backing for workflow, RBAC, CI/CD, civic AI")
bullets(s, [
    "Socio-spatial disparities in 311 smart-city governance — Sustainable Cities & Society (Sci. Direct, 2020)",
    "Ferraiolo & Kuhn — NIST model for role-based access control; RBAC on the Web (NIST publications)",
    "Hyun et al. — Jenkins automation impact on container deployment efficiency, Sensors 24(18), 2024",
    "Rai et al. — Automated full-stack deployment with CI/CD in DevOps, IJACTE 15(1), 2026; Sharma & Barua — DevSecOps CI/CD on AWS review, IJERT 2026",
], size=14, gap=0.72, bold_prefix=False)

# 17 — Thanks
s = prs.slides.add_slide(BLANK)
header(s, 17)
t = s.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(12.1), Inches(1.2))
tf = t.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Thank You!"
r.font.size = Pt(54); r.font.bold = True; r.font.color.rgb = NAVY
t2 = s.shapes.add_textbox(Inches(1.3), Inches(3.6), Inches(10.7), Inches(1.2))
tf = t2.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "SafeCity — Report it. Track it. Resolve it.\nLive demo: seeded Docker stack (18081 / 5174) + seeded accounts"
r.font.size = Pt(18); r.font.color.rgb = TEAL

prs.save(r"C:\projects\safecity0\SafeCity_PPT_v2.pptx")
print("saved SafeCity_PPT_v2.pptx, slides:", len(prs.slides))
